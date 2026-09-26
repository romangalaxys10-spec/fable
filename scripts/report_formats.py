#!/usr/bin/env python3
"""fable report formats — XLSX / DOCX / PPTX renderings of an audit run.

Reads the same secmon run JSONs as report.py and emits office formats:

  XLSX  findings tracker (openpyxl if installed, else SpreadsheetML —
        Excel-compatible XML, stdlib-only, always works)
  DOCX  narrative report (python-docx if installed, else skipped with note)
  PPTX  executive deck (python-pptx if installed, else skipped with note)

Usage:
  python3 report_formats.py <target> [--run-file run-XXXX.json] [--formats xlsx,docx,pptx] [--out-dir DIR]
"""

import datetime
import json
import os
import re
import sys

PLUGIN_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
BASE = os.path.expanduser("~/.fable/secmon")
ORDER = ["critical", "high", "medium", "low", "info"]


def find_run(target, run_file=None):
    slug = re.sub(r"[^a-z0-9.-]+", "-", target.lower().replace("https://", "").replace("http://", "")).strip("-")
    d = os.path.join(BASE, slug)
    name = run_file or (sorted(f for f in os.listdir(d) if f.startswith("run-"))[-1]
                        if os.path.isdir(d) and any(f.startswith("run-") for f in os.listdir(d)) else None)
    if not name:
        raise FileNotFoundError(f"no secmon runs for {target}")
    return json.load(open(os.path.join(d, name))), name


def sort_findings(findings):
    return sorted(findings, key=lambda f: ORDER.index(f.get("severity", "info")))


# ---------------------------------------------------------------- XLSX

def emit_xlsx(run, path):
    findings = sort_findings(run.get("findings", []))
    try:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Findings"
        ws.append(["severity", "id", "title", "detail/recommendation", "uid"])
        for f in findings:
            ws.append([f.get("severity", ""), f.get("id", ""),
                       f.get("title", ""), f.get("recommendation", ""), f.get("uid", "")])
        widths = [10, 22, 46, 70, 14]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + i)].width = w
        ws2 = wb.create_sheet("Summary")
        for k, v in (run.get("summary") or {}).items():
            ws2.append([k, v])
        wb.save(path)
        return "xlsx (openpyxl)"
    except ImportError:
        pass
    # SpreadsheetML fallback: Excel-compatible XML, stdlib only
    import xml.sax.saxutils as sx
    rows = "".join(
        f"<Row><Cell><Data ss:Type=\"String\">{sx.escape(f.get('severity',''))}</Data></Cell>"
        f"<Cell><Data ss:Type=\"String\">{sx.escape(f.get('id',''))}</Data></Cell>"
        f"<Cell><Data ss:Type=\"String\">{sx.escape(f.get('title',''))}</Data></Cell>"
        f"<Cell><Data ss:Type=\"String\">{sx.escape(f.get('recommendation',''))}</Data></Cell></Row>"
        for f in findings)
    xml = f"""<?xml version="1.0"?>
<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet"
 xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet">
 <Worksheet ss:Name="Findings"><Table>{rows}</Table></Worksheet>
</Workbook>"""
    with open(os.path.splitext(path)[0] + ".xml", "w", encoding="utf-8") as fh:
        fh.write(xml)
    return "xlsx-fallback (SpreadsheetML .xml — open in Excel)"


# ---------------------------------------------------------------- DOCX

def emit_docx(run, path):
    try:
        from docx import Document
    except ImportError:
        return "docx skipped (python-docx not installed — pip install python-docx)"
    doc = Document()
    doc.add_heading(f"Fable Audit Report — {run.get('target','')}", 0)
    doc.add_paragraph(f"Scanned: {run.get('ts','')[:19]}")
    doc.add_heading("Summary", level=1)
    for k, v in (run.get("summary") or {}).items():
        doc.add_paragraph(f"{k}: {v}", style="List Bullet")
    doc.add_heading("Findings", level=1)
    for f in sort_findings(run.get("findings", [])):
        doc.add_heading(f"[{f.get('severity','').upper()}] {f.get('title','')}", level=2)
        if f.get("recommendation"):
            doc.add_paragraph(f"Fix: {f['recommendation']}")
    doc.save(path)
    return "docx (python-docx)"


# ---------------------------------------------------------------- PPTX

def emit_pptx(run, path):
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError:
        return "pptx skipped (python-pptx not installed — pip install python-pptx)"
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = f"Site Health — {run.get('target','')}"
    s1.placeholders[1].text = f"Fable audit · {run.get('ts','')[:10]}"

    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Summary"
    body = s2.placeholders[1].text_frame
    for k, v in (run.get("summary") or {}).items():
        body.add_paragraph().text = f"{k}: {v}"

    top = sort_findings(run.get("findings", []))[:5]
    if top:
        s3 = prs.slides.add_slide(prs.slide_layouts[1])
        s3.shapes.title.text = "Top findings"
        body = s3.placeholders[1].text_frame
        for f in top:
            body.add_paragraph().text = f"[{f.get('severity','')}] {f.get('title','')}"

    prs.save(path)
    return "pptx (python-pptx)"


def main():
    a = sys.argv[1:]
    if not a:
        print(json.dumps({"error": 'usage: report_formats.py <target> [--run-file …] [--formats xlsx,docx,pptx] [--out-dir DIR]'}))
        sys.exit(3)
    target = a[0]
    run_file = a[a.index("--run-file") + 1] if "--run-file" in a else None
    fmts = [x.strip() for x in (a[a.index("--formats") + 1].split(",") if "--formats" in a
                                else ["xlsx", "docx", "pptx"])]
    out_dir = a[a.index("--out-dir") + 1] if "--out-dir" in a else "."
    os.makedirs(out_dir, exist_ok=True)

    run, run_name = find_run(target, run_file)
    safe = re.sub(r"[^a-z0-9.-]+", "-", target.lower().replace("https://", "").replace("http://", "")).strip("-")
    stamp = datetime.datetime.now().strftime("%Y-%m-%d")
    emitted = {}
    for fmt in fmts:
        path = os.path.join(out_dir, f"fable-{fmt}-{safe}-{stamp}.{fmt if fmt != 'xlsx' else 'xlsx'}")
        try:
            if fmt == "xlsx":
                emitted["xlsx"] = emit_xlsx(run, path)
            elif fmt == "docx":
                emitted["docx"] = emit_docx(run, path)
            elif fmt == "pptx":
                emitted["pptx"] = emit_pptx(run, path)
        except Exception as e:
            emitted[fmt] = f"failed: {str(e)[:120]}"
    print(json.dumps({"run": run_name, "formats": emitted, "out_dir": out_dir}, indent=2))


if __name__ == "__main__":
    main()
